"""Tailored cover-letter generation (Jinja2 -> Markdown, DOCX, PDF)."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt
from jinja2 import Environment, FileSystemLoader, select_autoescape

from .config import COVERLETTERS_DIR, TEMPLATES_DIR, get_settings
from .jobs import JobDescription
from .logger import get_logger
from .pdf import render_text_pdf
from .resume import ResumeData
from .utils import slugify, unique_terms

log = get_logger("coverletter")

YEARS_EXPERIENCE = 6
_NUMERIC = re.compile(r"\d")


def _to_sentence(items: List[str]) -> str:
    items = [i for i in items if i]
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " and " + items[-1]


def build_highlights(job: JobDescription, resume: ResumeData, limit: int = 4) -> List[str]:
    """Pick the strongest, most relevant achievement bullets for *job*.

    Prefers quantified bullets (containing numbers) from experience sections,
    ranking those that mention JD keywords first.
    """
    matched = [m.lower() for m in job.matched_skills]
    candidates: List[str] = []
    for section in resume.sections:
        if not any(k in section.title.lower() for k in ("experience", "consulting", "work")):
            continue
        for block in section.blocks:
            if block.kind == "bullet" and _NUMERIC.search(block.text):
                candidates.append(block.text)

    def score(text: str) -> int:
        low = text.lower()
        return sum(1 for m in matched if m in low)

    candidates.sort(key=score, reverse=True)
    if not candidates:  # fallback: any bullets at all
        for section in resume.sections:
            for block in section.blocks:
                if block.kind == "bullet":
                    candidates.append(block.text)
    return unique_terms(candidates)[:limit]


def build_context(job: JobDescription, resume: ResumeData) -> Dict[str, Any]:
    """Assemble the Jinja2 render context shared by cover letter + email."""
    settings = get_settings()
    matched = unique_terms(job.matched_skills)
    company_clean = "" if job.company == "Unknown Company" else job.company

    opening = (
        f"Your posting for the {job.role} role stood out to me because it sits right "
        "in the space where I do my best work — taking a complex product problem and "
        "shaping it into clear flows, a consistent interface, and a design system a team "
        "can actually build on."
    )

    highlights = build_highlights(job, resume)
    highlights_block = "\n".join(f"- {h}" for h in highlights)
    matched_sentence = _to_sentence(matched[:6])

    # Pre-compute the optional "matched skills" paragraphs so the templates stay
    # free of Jinja whitespace-control quirks.
    if matched:
        cover_matched_paragraph = (
            "\nThe requirements you listed overlap closely with my day-to-day toolkit — "
            f"including {matched_sentence}. I am comfortable owning work end to end, from "
            "problem framing and user flows through high-fidelity UI, prototypes, design "
            "systems, and developer handoff.\n"
        )
        email_matched_line = (
            "\nMy experience lines up closely with the role — particularly "
            f"{matched_sentence}.\n"
        )
    else:
        cover_matched_paragraph = ""
        email_matched_line = ""

    return {
        "candidate_name": settings.candidate_name,
        "candidate_email": settings.candidate_email,
        "candidate_phone": settings.candidate_phone,
        "candidate_location": settings.candidate_location,
        "candidate_portfolio": settings.candidate_portfolio,
        "candidate_linkedin": settings.candidate_linkedin,
        "date": date.today().strftime("%B %d, %Y"),
        "company": company_clean or job.company,
        "role": job.role,
        "years_experience": YEARS_EXPERIENCE,
        "opening_paragraph": opening,
        "highlights": highlights,
        "highlights_block": highlights_block,
        "matched_skills": matched,
        "matched_skills_sentence": matched_sentence,
        "cover_matched_paragraph": cover_matched_paragraph,
        "email_matched_line": email_matched_line,
    }


class CoverLetterBuilder:
    """Render and persist a tailored cover letter in Markdown, DOCX and PDF."""

    def __init__(self, templates_dir: Path | None = None) -> None:
        self.templates_dir = templates_dir or TEMPLATES_DIR
        self.env = Environment(
            loader=FileSystemLoader(str(self.templates_dir)),
            autoescape=select_autoescape(enabled_extensions=(), default=False),
            trim_blocks=True,
            lstrip_blocks=True,
            keep_trailing_newline=True,
        )

    def render(self, job: JobDescription, resume: ResumeData) -> str:
        template = self.env.get_template("cover_letter.md")
        text = template.render(**build_context(job, resume))
        # Collapse the runs of blank lines Jinja conditionals can leave behind.
        return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"

    def write_docx(self, text: str, out_path: Path) -> Path:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)
        for line in text.split("\n"):
            stripped = line.strip()
            if stripped.startswith("- "):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)
        doc.save(str(out_path))
        log.info("Wrote cover letter DOCX -> %s", out_path)
        return out_path

    def build(self, job: JobDescription, resume: ResumeData) -> Dict[str, Path]:
        """Produce .md, .docx and .pdf cover letters. Returns their paths."""
        slug = slugify(job.company)
        text = self.render(job, resume)

        md_path = COVERLETTERS_DIR / f"{slug}_cover_letter.md"
        md_path.parent.mkdir(parents=True, exist_ok=True)
        md_path.write_text(text, encoding="utf-8")
        log.info("Wrote cover letter Markdown -> %s", md_path)

        docx_path = self.write_docx(text, COVERLETTERS_DIR / f"{slug}_cover_letter.docx")
        pdf_path = render_text_pdf(text, COVERLETTERS_DIR / f"{slug}_cover_letter.pdf")

        return {"md": md_path, "docx": docx_path, "pdf": pdf_path}
