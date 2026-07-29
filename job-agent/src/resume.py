"""ATS-optimised resume tailoring.

The master resume (``master_resume.docx``) is treated as strictly read-only.
:class:`ResumeBuilder` parses it into a structured model, tailors a copy toward a
specific :class:`~src.jobs.JobDescription` (surfacing matched keywords, adding an
ATS keyword block, injecting a role-specific summary line — never fabricating
experience), then writes ``generated/resumes/<company>_resume.docx``.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Literal

from docx import Document
from docx.shared import Pt, RGBColor

from .config import MASTER_RESUME, RESUMES_DIR
from .jobs import JobDescription
from .logger import get_logger
from .utils import slugify, unique_terms

log = get_logger("resume")

BlockKind = Literal["para", "bullet", "subhead"]


@dataclass
class Block:
    """A single line of resume content."""

    kind: BlockKind
    text: str


@dataclass
class Section:
    """A titled resume section holding ordered blocks."""

    title: str
    blocks: List[Block] = field(default_factory=list)


@dataclass
class ResumeData:
    """Structured, render-target-agnostic resume content."""

    name: str = ""
    headline: str = ""
    contact_lines: List[str] = field(default_factory=list)
    sections: List[Section] = field(default_factory=list)


class ResumeBuilder:
    """Parse the master resume and produce tailored copies."""

    def __init__(self, master_path: Path | None = None) -> None:
        self.master_path = master_path or MASTER_RESUME
        if not self.master_path.exists():
            raise FileNotFoundError(
                f"Master resume not found at {self.master_path}. "
                "Run `python scripts/seed_master_resume.py` to create it."
            )

    # --- Parsing ---------------------------------------------------------
    def load_master(self) -> ResumeData:
        """Parse ``master_resume.docx`` into a :class:`ResumeData` (read-only)."""
        doc = Document(str(self.master_path))
        data = ResumeData()
        current: Section | None = None
        seen_first_section = False

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()

            if style == "title" and not data.name:
                data.name = text
            elif style == "subtitle" and not data.headline:
                data.headline = text
            elif style.startswith("heading 1"):
                current = Section(title=text)
                data.sections.append(current)
                seen_first_section = True
            elif style.startswith("heading 2"):
                if current is not None:
                    current.blocks.append(Block("subhead", text))
            elif style.startswith("list"):
                if current is not None:
                    current.blocks.append(Block("bullet", text))
            else:  # normal paragraph
                if not seen_first_section:
                    data.contact_lines.append(text)
                elif current is not None:
                    current.blocks.append(Block("para", text))

        log.info(
            "Loaded master resume: %d sections, name=%r",
            len(data.sections), data.name,
        )
        return data

    # --- Tailoring -------------------------------------------------------
    def tailor(self, job: JobDescription) -> ResumeData:
        """Return a tailored copy of the master resume for *job*.

        Tailoring is additive and truthful:
          * a role-specific summary line is prepended to the summary,
          * a "Key Skills" ATS section (matched keywords) is inserted near the top,
          * core-capability bullets that match the JD are moved to the front.
        No experience is invented or removed.
        """
        data = self.load_master()
        matched = job.matched_skills

        # 1) Role-aware summary line.
        for section in data.sections:
            if "summary" in section.title.lower():
                lead = (
                    f"Targeting the {job.role} role"
                    + (f" at {job.company}" if job.company and job.company != "Unknown Company" else "")
                    + "."
                )
                if matched:
                    lead += " Directly relevant strengths include " + _to_sentence(matched[:6]) + "."
                section.blocks.insert(0, Block("para", lead))
                break

        # 2) ATS keyword section right after the summary.
        if matched:
            ats = Section(
                title="Key Skills",
                blocks=[Block("para", " · ".join(unique_terms(matched)))],
            )
            insert_at = 1 if data.sections else 0
            data.sections.insert(insert_at, ats)

        # 3) Reorder Core Capabilities to surface matched bullets first.
        for section in data.sections:
            if "capabilit" in section.title.lower():
                section.blocks.sort(
                    key=lambda b: 0 if _bullet_matches(b, matched) else 1
                )
                break

        return data

    # --- Rendering -------------------------------------------------------
    def write_docx(self, data: ResumeData, out_path: Path) -> Path:
        """Render *data* to a .docx file at *out_path*."""
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10.5)

        if data.name:
            p = doc.add_paragraph(data.name, style="Title")
        if data.headline:
            doc.add_paragraph(data.headline, style="Subtitle")
        for line in data.contact_lines:
            doc.add_paragraph(line)

        for section in data.sections:
            heading = doc.add_paragraph(section.title, style="Heading 1")
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            for block in section.blocks:
                if block.kind == "subhead":
                    doc.add_paragraph(block.text, style="Heading 2")
                elif block.kind == "bullet":
                    doc.add_paragraph(block.text, style="List Bullet")
                else:
                    doc.add_paragraph(block.text)

        doc.save(str(out_path))
        log.info("Wrote tailored resume DOCX -> %s", out_path)
        return out_path

    def build(self, job: JobDescription) -> tuple[ResumeData, Path]:
        """Tailor + write the resume docx. Returns (data, docx_path)."""
        data = self.tailor(job)
        slug = slugify(job.company)
        out_path = RESUMES_DIR / f"{slug}_resume.docx"
        self.write_docx(data, out_path)
        return data, out_path


def _bullet_matches(block: Block, matched: List[str]) -> bool:
    low = block.text.lower()
    return any(m.lower() in low for m in matched)


def _to_sentence(items: List[str]) -> str:
    """Join a list into an English clause: 'a, b and c'."""
    items = [i for i in items if i]
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " and " + items[-1]
