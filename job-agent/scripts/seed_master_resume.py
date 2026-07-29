"""One-off generator for ``master_resume.docx`` and ``templates/resume_template.docx``.

The agent never edits the master resume at runtime; this script simply creates a
well-structured starting document (styled with Title / Heading 1 / Heading 2 /
List Bullet paragraphs) that :mod:`src.resume` knows how to parse and tailor.

Run:  python scripts/seed_master_resume.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

# Allow running as a standalone script.
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

NAME = "Mohd Hayaat Ali"
HEADLINE = "Product Designer | UI/UX, Interaction and Visual Systems"
CONTACT = [
    "mohdhayaat1@outlook.com | +91-7905194153 | Delhi NCR, India",
    "Portfolio: https://workofhayaat.framer.website | Open to relocation and remote opportunities",
]

# Each section: (title, [("para"|"bullet"|"subhead", text), ...])
SECTIONS = [
    (
        "Professional Summary",
        [
            (
                "para",
                "Product and UI/UX designer with 6+ years of experience across shipped "
                "mobile apps, web products, B2B platforms, internal tools, healthcare, "
                "fuel-tech, logistics, fintech, AI automation, e-commerce, and brand-led "
                "digital experiences. I work from problem framing and user flows through "
                "wireframes, high-fidelity UI, prototypes, design systems, usability "
                "testing, and developer handoff. My strongest work simplifies complex "
                "operational journeys while maintaining clear interaction patterns and "
                "polished visual execution.",
            ),
        ],
    ),
    (
        "Core Capabilities",
        [
            ("bullet", "Product discovery, user flows, information architecture, and journey mapping"),
            ("bullet", "Wireframing, interaction design, responsive web and mobile design, and prototyping"),
            ("bullet", "Usability testing, A/B testing, and analytics-driven iteration"),
            ("bullet", "Design systems, Figma components, Auto Layout, and component documentation"),
            ("bullet", "Accessibility awareness, developer handoff, design QA, and stakeholder presentations"),
            ("bullet", "B2B SaaS, operational dashboards, payment and wallet UX, KYC, logistics and fleet workflows, healthcare journeys, and marketplaces"),
        ],
    ),
    (
        "Professional Experience",
        [
            ("subhead", "Creative Designer | Crevia — Delhi NCR, India | June 2026 - Present"),
            ("bullet", "Create graphics, videos, social content, product imagery, website assets, and campaign material across brands including Mymy, Al Yamin, Klay Consultants, and ATK Shoe Manufacturer."),
            ("bullet", "Designed and built the Mymy perfume-brand website in Wix, covering responsive page structure, product presentation, visual direction, and launch-ready assets."),
            ("bullet", "Edit podcast videos, product images, and campaign content while maintaining consistent visual systems across web and social touchpoints."),

            ("subhead", "Product Designer | I-DOD — New Delhi, India | July 2025 - March 2026"),
            ("bullet", "Designed onboarding, KYC verification, profile creation, matching journeys, and beta feedback loops for an early-stage relationship platform."),
            ("bullet", "Built the design system from scratch with reusable Figma components, interaction patterns, states, and implementation-ready documentation."),
            ("bullet", "Worked directly with founders and developers to turn ambiguous product requirements into user flows, prototypes, and high-fidelity interfaces."),

            ("subhead", "Product Designer | FuelBuddy — Gurgaon, India | July 2023 - June 2024"),
            ("bullet", "Redesigned diesel ordering from a 20-22 step journey into a focused location, quantity, schedule, and payment flow, improving completion from 62% to 78%."),
            ("bullet", "Reduced payment transaction errors by 38% through wallet improvements, delegated access, secondary-user controls, and configurable spend limits."),
            ("bullet", "Designed customer, FMS, pilot, and driver experiences across India and UAE, spanning mobile apps, B2B web products, analytics, maps, live tracking, and asset management."),
            ("bullet", "Reduced manual support input by 43% and improved ticket handling and resolution by 57% through structured support and ticketing workflows."),
            ("bullet", "Built the FuelBuddy design system from scratch across consumer, B2B, franchise, field, and operational products."),

            ("subhead", "Product Designer | Uncover by Meddo — Gurgaon, India | March 2022 - May 2023"),
            ("bullet", "Reduced appointment booking from six steps to four, improving completion from 71% to 83% over eight weeks using analytics and user feedback."),
            ("bullet", "Redesigned doctor profiles, increasing profile views by 28% and appointment requests by 15% in a 5,000-user A/B test."),
            ("bullet", "Designed patient, doctor, and back-office experiences spanning discovery, appointments, healthcare memberships, lab tests, and digital records."),
            ("bullet", "Conducted usability testing with 12 patients and 8 doctors, translating findings into product, design-system, and engineering priorities."),

            ("subhead", "UI Designer | AcadPlaza — Remote | June 2020 - March 2022"),
            ("bullet", "Redesigned course catalog and search for a learning marketplace, increasing enrollments by 18% quarter over quarter."),
            ("bullet", "Created responsive web and mobile interfaces, wireframes, component patterns, and discovery journeys with an eight-person product, engineering, and content team."),
        ],
    ),
    (
        "Selected Consulting and Product Work",
        [
            ("bullet", "Vgen23: Designed a genetic interpretation and reporting web application that simplifies clinical and genomic workflows for geneticists."),
            ("bullet", "TS Logix Peru: Designed the public website and an internal WMS for logistics, inventory, warehousing, and pharmaceutical distribution workflows."),
            ("bullet", "Maximor AI: Designed a conversion-focused CFO offer page for an AI finance-automation platform covering close, reconciliation, reporting, and forecasting."),
            ("bullet", "Kama Capital: Designed a multi-asset trading website covering forex, commodities, indices, stocks, account types, and onboarding journeys."),
        ],
    ),
    (
        "Tools",
        [
            ("para", "Figma, FigJam, Framer, Wix, Adobe Photoshop, Illustrator, After Effects, Rive, Miro, Adobe XD, Google Analytics, Hotjar"),
        ],
    ),
    (
        "Education",
        [
            ("para", "BBA, Business Administration — Sam Higginbottom University of Agriculture, Technology and Sciences | 2017 - 2020"),
        ],
    ),
]


def _add_styled(doc: Document) -> None:
    doc.add_paragraph(NAME, style="Title")
    doc.add_paragraph(HEADLINE, style="Subtitle")
    for line in CONTACT:
        doc.add_paragraph(line)
    for title, blocks in SECTIONS:
        doc.add_paragraph(title, style="Heading 1")
        for kind, text in blocks:
            if kind == "subhead":
                doc.add_paragraph(text, style="Heading 2")
            elif kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)


def build_master(path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    _add_styled(doc)
    doc.save(str(path))
    print(f"Wrote {path}")


def build_template(path: Path) -> None:
    """A minimal style-base document used as the starting point for tailored resumes."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.save(str(path))
    print(f"Wrote {path}")


def main() -> None:
    build_master(ROOT / "master_resume.docx")
    (ROOT / "templates").mkdir(parents=True, exist_ok=True)
    build_template(ROOT / "templates" / "resume_template.docx")


if __name__ == "__main__":
    main()
