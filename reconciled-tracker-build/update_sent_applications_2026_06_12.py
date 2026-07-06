from pathlib import Path

import openpyxl

TRACKERS = [Path("reconciled-tracker-build/master_job_tracker_reconciled.xlsx")]
SIMPLE_TRACKER = Path("reconciled-tracker-build/current_tracker.xlsx")

ROWS = [
    {
        "Country / Region": "India",
        "Type": "Direct Employer",
        "Company / Agency": "Godrej Industries Group - DISCO",
        "Company Description": "In-house Design, Innovation, Strategy & Creative Collaboration studio",
        "Contact Person": "Rominder Sapra",
        "Email Sent To": "rominder.sapra@godrejinds.com",
        "Email / Contact": "rominder.sapra@godrejinds.com",
        "Role Applied For": "UI & Interaction Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "High",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Direct email",
        "Notes": "Sent tailored DISCO resume. Gmail message ID: 19ebbbd1041508a9. Mumbai full-time role.",
    },
    {
        "Country / Region": "India",
        "Type": "Direct Employer",
        "Company / Agency": "Appinventiv Technologies",
        "Company Description": "Technology services company hiring UI/UX Designer",
        "Contact Person": "Shivangi Shukla",
        "Email Sent To": "shivangi.shukla@appinventiv.com",
        "Email / Contact": "shivangi.shukla@appinventiv.com",
        "Role Applied For": "UI/UX Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "High",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Naukri mailer / direct email",
        "Notes": "Sent tailored Appinventiv resume. Gmail message ID: 19ebbbfe8acd2caa. Noida in-office role.",
    },
    {
        "Country / Region": "India",
        "Type": "Direct Employer",
        "Company / Agency": "Kukami Technology",
        "Company Description": "Technology company hiring UI/UX Designer",
        "Contact Person": "HR",
        "Email Sent To": "hrd.kukamitechnology@gmail.com",
        "Email / Contact": "hrd.kukamitechnology@gmail.com",
        "Phone / WhatsApp": "9998477789",
        "Role Applied For": "UI/UX Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "High",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Direct email",
        "Notes": "Sent tailored Kukami resume. Gmail message ID: 19ebbcd61584a834. Surat on-site role.",
    },
    {
        "Country / Region": "India",
        "Type": "Direct Employer",
        "Company / Agency": "Monotype",
        "Company Description": "Product company hiring for Lead Product Designer scope",
        "Contact Person": "Nidhi Gupta",
        "Email Sent To": "nidhi.gupta@monotype.com",
        "Email / Contact": "nidhi.gupta@monotype.com",
        "Role Applied For": "Lead Product Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "Medium",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Direct email",
        "Notes": "Sent tailored Monotype resume with explicit 4+ years framing. Gmail message ID: 19ebbd921ed10d29.",
    },
    {
        "Country / Region": "India",
        "Type": "Direct Employer",
        "Company / Agency": "Foldr",
        "Company Description": "Remote full-time UI/UX and identity design role",
        "Contact Person": "Hiring Team",
        "Email Sent To": "people@foldr.studio",
        "Email / Contact": "people@foldr.studio",
        "Role Applied For": "UI/UX & Identity Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "High",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Direct email",
        "Notes": "Sent tailored Foldr resume. Gmail message ID: 19ebcad0882171da. Remote full-time role.",
        "Last Gmail Thread ID": "19ebcad0882171da",
        "Last Gmail Message ID": "19ebcad0882171da",
    },
    {
        "Country / Region": "UAE / Dubai",
        "Type": "Agency",
        "Company / Agency": "Inspire Selection",
        "Company Description": "Recruitment agency",
        "Contact Person": "Daniel Asare",
        "Email Sent To": "info@inspireselection.com",
        "Email / Contact": "info@inspireselection.com",
        "Role Applied For": "UI/UX Designer / Product Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "High",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Agency email",
        "Notes": "Sent tailored agency resume. Gmail message ID: 19ebcb58036e1ca4.",
        "Last Gmail Thread ID": "19ebcb58036e1ca4",
        "Last Gmail Message ID": "19ebcb58036e1ca4",
    },
    {
        "Country / Region": "UAE / Dubai",
        "Type": "Agency",
        "Company / Agency": "360agency",
        "Company Description": "Agency / recruiter",
        "Contact Person": "Recruitment Team",
        "Email Sent To": "apply@360agency.me",
        "Email / Contact": "apply@360agency.me",
        "Role Applied For": "Product Designer / UI/UX Designer",
        "Status": "Applied",
        "Date Email Sent": "2026-06-12",
        "Date Applied": "2026-06-12",
        "Reply Status": "No Reply",
        "Confidence": "Medium",
        "Needs Manual Review": "No",
        "Recommended Action": "Follow up if no reply",
        "Source": "Agency email",
        "Notes": "Sent UAE agencies resume. Gmail message ID: 19ebcbd4707cfbfc.",
        "Last Gmail Thread ID": "19ebcbd4707cfbfc",
        "Last Gmail Message ID": "19ebcbd4707cfbfc",
    },
]


def headers(ws):
    return {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}


def row_exists(ws, header_map, row):
    company_col = header_map["Company / Agency"]
    email_col = header_map["Email / Contact"]
    company = row["Company / Agency"].lower()
    email = row["Email / Contact"].lower()
    for idx in range(2, ws.max_row + 1):
        existing_company = str(ws.cell(idx, company_col).value or "").lower()
        existing_email = str(ws.cell(idx, email_col).value or "").lower()
        if company in existing_company or email == existing_email:
            return idx
    return None


def upsert_row(ws, row):
    header_map = headers(ws)
    target = row_exists(ws, header_map, row) or ws.max_row + 1
    for key, value in row.items():
        col = header_map.get(key)
        if col:
            ws.cell(target, col).value = value
    return target


for tracker in TRACKERS:
    wb = openpyxl.load_workbook(tracker)
    updated = []
    for sheet_name in ["Master Leads", "Direct Employers"]:
        ws = wb[sheet_name]
        for row in ROWS:
            updated.append((sheet_name, row["Company / Agency"], upsert_row(ws, row)))
    wb.save(tracker)
    print(tracker)
    for item in updated:
        print(item)

wb = openpyxl.load_workbook(SIMPLE_TRACKER)
ws = wb["Master Job Tracker"]
header_map = headers(ws)
updated = []
for row in ROWS:
    simple = {
        "Country / Region": row["Country / Region"],
        "Company / Agency": row["Company / Agency"],
        "Email": row["Email / Contact"],
        "Type": row["Type"],
        "Status": row["Status"],
        "Date Applied": row["Date Applied"],
        "Follow-up Date": "",
        "Notes": row["Notes"],
        "Opening Scan Status": "Applied via email",
        "Matched Role(s)": row["Role Applied For"],
        "Opening Source URL(s)": "",
        "Confidence": row["Confidence"],
        "Recommended Action": row["Recommended Action"],
        "Scan Date": "2026-06-12",
    }
    target = row_exists(ws, {"Company / Agency": header_map["Company / Agency"], "Email / Contact": header_map["Email"]}, {
        "Company / Agency": row["Company / Agency"],
        "Email / Contact": row["Email / Contact"],
    }) or ws.max_row + 1
    for key, value in simple.items():
        ws.cell(target, header_map[key]).value = value
    updated.append(("Master Job Tracker", row["Company / Agency"], target))
wb.save(SIMPLE_TRACKER)
print(SIMPLE_TRACKER)
for item in updated:
    print(item)
